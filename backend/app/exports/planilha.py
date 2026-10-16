"""Geração de planilhas Excel (openpyxl) — etapa 10 (RF-012/013).

Gera arquivos ``.xlsx`` a partir de dados **já validados** (nunca de texto bruto
do OCR). Regras aplicadas:

- Só entram registros exportáveis (``app.business.exportavel``): confirmados ou
  corrigidos e validados. Dúvidas/pendências nunca aparecem como confirmadas.
- Nenhum horário é inventado: batidas ausentes ficam em branco e dias sem
  marcação reconhecidos ficam sem total.
- Organização por funcionário e dia; datas, batidas e totais calculados pelas
  regras existentes (``app.business.calculo``).
- Textos que o Excel interpretaria como fórmula são neutralizados
  (CONTEXT §8, camada 5).
- A coluna "Hora Extra" não existe no modelo e **não** é exportada.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..business.calculo import calcular_horas
from ..business.exportavel import registros_exportaveis
from ..business.horarios import formatar_duracao
from ..persistence.models import FichaPonto, RegistroDiario

#: Prefixos que o Excel pode interpretar como início de fórmula.
PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")

#: Cabeçalhos das linhas por dia (a coluna "Hora Extra" é intencionalmente ausente).
CABECALHOS = (
    "Dia",
    "Data",
    "Dia da semana",
    "Entrada",
    "Saída almoço",
    "Volta almoço",
    "Saída",
    "Total",
    "Observação",
)

_FONTE_TITULO = Font(bold=True, size=14)
_FONTE_NEGRITO = Font(bold=True)


def celula_segura(valor: Any) -> Any:
    """Neutraliza textos que o Excel interpretaria como fórmula."""
    if isinstance(valor, str) and valor[:1] in PREFIXOS_FORMULA:
        return "'" + valor
    return valor


def _append(linha: Worksheet, valores: list[Any]) -> None:
    linha.append([celula_segura(valor) for valor in valores])


def _referencia(ficha: FichaPonto) -> str | None:
    if ficha.mes is None or ficha.ano is None:
        return None
    return f"{ficha.mes:02d}/{ficha.ano}"


def _data(ficha: FichaPonto, registro: RegistroDiario) -> str | None:
    if registro.dia is None or ficha.mes is None or ficha.ano is None:
        return None
    return f"{registro.dia:02d}/{ficha.mes:02d}/{ficha.ano}"


def _funcionario(ficha: FichaPonto) -> str:
    return ficha.funcionario.nome or ficha.funcionario.codigo or "—"


def total_registro(registro: RegistroDiario) -> str | None:
    """Total confirmado (``HH:MM``) do registro, ou ``None`` se não confirmado."""
    calculo = calcular_horas(registro)
    return calculo.total_formatado() if calculo.confirmado else None


def _registros_ordenados(ficha: FichaPonto) -> list[RegistroDiario]:
    return sorted(
        registros_exportaveis(ficha), key=lambda r: (r.dia is None, r.dia or 0)
    )


def _minutos_confirmados(registro: RegistroDiario) -> int:
    calculo = calcular_horas(registro)
    return calculo.total_minutos or 0 if calculo.confirmado else 0


def _ajustar_larguras(ws: Worksheet, minimo: int = 12, maximo: int = 40) -> None:
    for coluna in ws.iter_cols():
        letra = get_column_letter(coluna[0].column)
        maior = max(
            (len(str(celula.value)) for celula in coluna if celula.value is not None),
            default=0,
        )
        ws.column_dimensions[letra].width = min(max(maior + 2, minimo), maximo)


def gerar_planilha_funcionario(ficha: FichaPonto, destino: Path) -> int:
    """Gera o ``.xlsx`` de uma ficha (funcionário + período) e retorna o total de linhas."""
    registros = _registros_ordenados(ficha)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ponto"

    _append(ws, ["PontoFácil — Folha de ponto"])
    ws["A1"].font = _FONTE_TITULO
    _append(ws, ["Funcionário:", _funcionario(ficha)])
    if ficha.funcionario.codigo:
        _append(ws, ["Código:", ficha.funcionario.codigo])
    if ficha.funcionario.empresa:
        _append(ws, ["Empresa:", ficha.funcionario.empresa])
    _append(ws, ["Período:", _referencia(ficha) or "não identificado"])
    _append(ws, [""])

    _append(ws, list(CABECALHOS))
    for celula in ws[ws.max_row]:
        celula.font = _FONTE_NEGRITO

    total_minutos = 0
    for registro in registros:
        total_minutos += _minutos_confirmados(registro)
        _append(
            ws,
            [
                registro.dia,
                _data(ficha, registro),
                registro.dia_semana,
                registro.entrada,
                registro.saida_almoco,
                registro.volta_almoco,
                registro.saida,
                total_registro(registro),
                registro.observacao,
            ],
        )

    if registros:
        _append(ws, [""])
        _append(
            ws,
            [
                "Total confirmado",
                "",
                "",
                "",
                "",
                "",
                "",
                formatar_duracao(total_minutos),
            ],
        )
        ws.cell(row=ws.max_row, column=1).font = _FONTE_NEGRITO

    _ajustar_larguras(ws)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return len(registros)


def gerar_planilha_consolidada(fichas: list[FichaPonto], destino: Path) -> int:
    """Gera o ``.xlsx`` consolidado (Resumo + Consolidado) e retorna as linhas exportadas."""
    wb = Workbook()

    resumo = wb.active
    resumo.title = "Resumo"
    _append(resumo, ["PontoFácil — Resumo consolidado"])
    resumo["A1"].font = _FONTE_TITULO
    _append(resumo, [""])
    _append(
        resumo,
        ["Funcionário", "Código", "Período", "Dias exportados", "Total confirmado"],
    )
    for celula in resumo[resumo.max_row]:
        celula.font = _FONTE_NEGRITO

    consolidado = wb.create_sheet("Consolidado")
    _append(consolidado, ["Funcionário", "Período", *CABECALHOS])
    for celula in consolidado[1]:
        celula.font = _FONTE_NEGRITO

    total_linhas = 0
    total_geral = 0
    for ficha in fichas:
        registros = _registros_ordenados(ficha)
        if not registros:
            continue
        minutos_ficha = 0
        for registro in registros:
            minutos_ficha += _minutos_confirmados(registro)
            _append(
                consolidado,
                [
                    _funcionario(ficha),
                    _referencia(ficha),
                    registro.dia,
                    _data(ficha, registro),
                    registro.dia_semana,
                    registro.entrada,
                    registro.saida_almoco,
                    registro.volta_almoco,
                    registro.saida,
                    total_registro(registro),
                    registro.observacao,
                ],
            )
            total_linhas += 1
        total_geral += minutos_ficha
        _append(
            resumo,
            [
                _funcionario(ficha),
                ficha.funcionario.codigo,
                _referencia(ficha) or "não identificado",
                len(registros),
                formatar_duracao(minutos_ficha),
            ],
        )

    _append(resumo, [""])
    _append(resumo, ["Total geral", "", "", total_linhas, formatar_duracao(total_geral)])
    resumo.cell(row=resumo.max_row, column=1).font = _FONTE_NEGRITO

    _ajustar_larguras(resumo)
    _ajustar_larguras(consolidado)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return total_linhas
