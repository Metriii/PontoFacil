"""Testes da geração de planilhas Excel (etapa 10, RF-012).

Verificam: inclusão apenas de registros exportáveis, totais pelas regras
existentes, ausência de horários inventados, neutralização de fórmulas e a
ausência da coluna "Hora Extra".
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.exports import (
    CABECALHOS,
    celula_segura,
    gerar_planilha_consolidada,
    gerar_planilha_funcionario,
    total_registro,
)
from app.persistence import FichaPonto, Funcionario, RegistroDiario

COMPLETO = {
    "entrada": "08:00",
    "saida_almoco": "12:00",
    "volta_almoco": "13:00",
    "saida": "17:00",
}


def _ficha(registros: list[RegistroDiario], *, nome="Maria", mes=8, ano=2026) -> FichaPonto:
    return FichaPonto(
        funcionario=Funcionario(funcionario_id="f1", nome=nome, codigo="007"),
        mes=mes,
        ano=ano,
        registros=registros,
    )


def _registro(registro_id: str, dia: int, **campos) -> RegistroDiario:
    return RegistroDiario(registro_id=registro_id, dia=dia, origem="a.png", **campos)


def _linhas(ws) -> list[list[object]]:
    return [[celula.value for celula in linha] for linha in ws.iter_rows()]


def test_celula_segura_neutraliza_formulas() -> None:
    assert celula_segura("=SUM(A1)") == "'=SUM(A1)"
    assert celula_segura("+1") == "'+1"
    assert celula_segura("@x") == "'@x"
    assert celula_segura("-5") == "'-5"
    assert celula_segura("08:00") == "08:00"
    assert celula_segura(12) == 12
    assert celula_segura(None) is None


def test_total_registro_usa_regras_existentes() -> None:
    completo = _registro("r1", 3, status_validacao="confirmado", **COMPLETO)
    assert total_registro(completo) == "08:00"

    incompleto = _registro(
        "r2", 4, status_validacao="pendente", entrada="08:00", saida="17:00"
    )
    assert total_registro(incompleto) is None


def test_planilha_funcionario_inclui_confirmados_e_total(tmp_path: Path) -> None:
    ficha = _ficha(
        [
            _registro("r1", 3, status_validacao="confirmado", dia_semana="segunda-feira", **COMPLETO),
            _registro("r2", 4, status_validacao="pendente", entrada="08:00", saida="17:00"),
        ]
    )
    destino = tmp_path / "func.xlsx"
    linhas_exportadas = gerar_planilha_funcionario(ficha, destino)

    assert linhas_exportadas == 1
    valores = _linhas(load_workbook(destino)["Ponto"])
    dados = [linha for linha in valores if linha and linha[0] == 3]
    assert len(dados) == 1
    assert dados[0][1] == "03/08/2026"
    assert dados[0][3:7] == ["08:00", "12:00", "13:00", "17:00"]
    assert dados[0][7] == "08:00"
    # dia 4 (pendente) não é exportado
    assert all(linha[0] != 4 for linha in valores if linha)
    # total confirmado presente
    assert any(linha and linha[0] == "Total confirmado" and linha[7] == "08:00" for linha in valores)


def test_planilha_exclui_invalidos_e_ilegiveis(tmp_path: Path) -> None:
    ficha = _ficha(
        [
            _registro("r1", 1, status_validacao="invalido", entrada="25:99"),
            _registro(
                "r2", 2, status_validacao="pendente", entrada="08:00", valor_bruto={"saida": "O8:OO"}
            ),
        ]
    )
    destino = tmp_path / "func.xlsx"
    assert gerar_planilha_funcionario(ficha, destino) == 0
    valores = _linhas(load_workbook(destino)["Ponto"])
    assert all(linha[0] not in (1, 2) for linha in valores if linha)


def test_planilha_inclui_corrigido_validado(tmp_path: Path) -> None:
    ficha = _ficha(
        [_registro("r1", 5, status_validacao="corrigido", **COMPLETO)]
    )
    destino = tmp_path / "func.xlsx"
    assert gerar_planilha_funcionario(ficha, destino) == 1


def test_planilha_exclui_corrigido_ainda_incompleto(tmp_path: Path) -> None:
    ficha = _ficha(
        [
            _registro(
                "r1", 6, status_validacao="corrigido", entrada="08:00", saida="17:00"
            )
        ]
    )
    destino = tmp_path / "func.xlsx"
    assert gerar_planilha_funcionario(ficha, destino) == 0


def test_dia_sem_marcacao_confirmado_entra_sem_total(tmp_path: Path) -> None:
    ficha = _ficha(
        [_registro("r1", 7, status_validacao="confirmado", tipo_dia="fim_de_semana")]
    )
    destino = tmp_path / "func.xlsx"
    assert gerar_planilha_funcionario(ficha, destino) == 1
    valores = _linhas(load_workbook(destino)["Ponto"])
    dados = [linha for linha in valores if linha and linha[0] == 7]
    assert dados[0][3:8] == [None, None, None, None, None]


def test_neutraliza_formula_no_nome_do_funcionario(tmp_path: Path) -> None:
    ficha = _ficha(
        [_registro("r1", 3, status_validacao="confirmado", **COMPLETO)],
        nome="=cmd|'/c calc'!A0",
    )
    destino = tmp_path / "func.xlsx"
    gerar_planilha_funcionario(ficha, destino)
    valores = _linhas(load_workbook(destino)["Ponto"])
    linha_funcionario = next(linha for linha in valores if linha and linha[0] == "Funcionário:")
    assert str(linha_funcionario[1]).startswith("'")


def test_cabecalho_nao_tem_hora_extra() -> None:
    assert all("Hora Extra" not in coluna for coluna in CABECALHOS)


def test_planilha_consolidada_agrega_e_resume(tmp_path: Path) -> None:
    fichas = [
        _ficha([_registro("r1", 3, status_validacao="confirmado", **COMPLETO)], nome="Maria"),
        _ficha([_registro("r2", 1, status_validacao="confirmado", **COMPLETO)], nome="João"),
    ]
    destino = tmp_path / "consolidado.xlsx"
    total = gerar_planilha_consolidada(fichas, destino)

    assert total == 2
    wb = load_workbook(destino)
    assert set(wb.sheetnames) == {"Resumo", "Consolidado"}
    resumo = _linhas(wb["Resumo"])
    assert any(linha and linha[0] == "Total geral" and linha[4] == "16:00" for linha in resumo)
    consolidado = _linhas(wb["Consolidado"])
    assert consolidado[0][:2] == ["Funcionário", "Período"]
    assert sum(1 for linha in consolidado[1:] if linha and isinstance(linha[2], int)) == 2
