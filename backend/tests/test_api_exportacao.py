"""Testes dos endpoints de exportação e download (etapas 10-11, RF-012/013/015)."""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from .conftest import REGISTRO_COMPLETO, enviar_lote, folha_com, montar_cliente


def _laudo(cliente, lote_id):
    return cliente.post(f"/api/lotes/{lote_id}/exportar")


def test_exportar_gera_planilhas_e_permite_download(tmp_path) -> None:
    cliente, _ = montar_cliente(tmp_path, folha_com([dict(REGISTRO_COMPLETO)]))
    lote_id = enviar_lote(cliente).json()["lote_id"]

    resposta = _laudo(cliente, lote_id)
    assert resposta.status_code == 200
    corpo = resposta.json()
    assert corpo["total_registros"] == 1
    assert len(corpo["planilhas"]) == 1
    assert corpo["consolidado"] is not None
    assert corpo["planilhas"][0]["arquivo"].endswith(".xlsx")

    download = cliente.get(corpo["planilhas"][0]["url"])
    assert download.status_code == 200
    assert download.content[:2] == b"PK"  # .xlsx é um arquivo ZIP
    assert "attachment" in download.headers["content-disposition"]

    assert cliente.get(corpo["consolidado"]["url"]).status_code == 200

    # registrados no lote e listados em /downloads
    status = cliente.get(f"/api/lotes/{lote_id}").json()
    assert len(status["arquivos_gerados"]) == 2
    listados = cliente.get(f"/api/lotes/{lote_id}/downloads").json()
    assert len(listados) == 2


def test_consolidado_so_inclui_exportaveis(tmp_path) -> None:
    extrator = folha_com([dict(REGISTRO_COMPLETO), {"entrada": "08:00", "saida": "17:00"}])
    cliente, _ = montar_cliente(tmp_path, extrator)
    lote_id = enviar_lote(cliente).json()["lote_id"]

    corpo = _laudo(cliente, lote_id).json()
    assert corpo["total_registros"] == 1

    bruto = cliente.get(corpo["consolidado"]["url"]).content
    wb = load_workbook(BytesIO(bruto))
    linhas = [
        [celula.value for celula in linha]
        for linha in wb["Consolidado"].iter_rows()
    ]
    dados = [linha for linha in linhas[1:] if linha and isinstance(linha[2], int)]
    assert len(dados) == 1
    assert dados[0][3] == "01/08/2026"


def test_exportar_inclui_corrigido_validado(tmp_path) -> None:
    extrator = folha_com(
        [
            {
                "entrada": "08:00",
                "saida_almoco": "12:00",
                "volta_almoco": "13:00",
                "saida": None,
                "valor_bruto": {"saida": "O8:OO"},
            }
        ]
    )
    cliente, _ = montar_cliente(tmp_path, extrator)
    lote_id = enviar_lote(cliente).json()["lote_id"]

    # Antes da correção não há registros exportáveis.
    assert _laudo(cliente, lote_id).json()["total_registros"] == 0

    registro = cliente.get(f"/api/lotes/{lote_id}/resultados").json()["fichas"][0][
        "registros"
    ][0]
    corrigido = cliente.patch(
        f"/api/registros/{registro['registro_id']}",
        json={"campo": "saida", "valor": "17:00"},
    )
    assert corrigido.status_code == 200
    assert corrigido.json()["status_validacao"] == "corrigido"

    corpo = _laudo(cliente, lote_id).json()
    assert corpo["total_registros"] == 1
    assert len(corpo["planilhas"]) == 1


def test_lote_sem_dados_exportaveis(tmp_path) -> None:
    cliente, _ = montar_cliente(
        tmp_path, folha_com([{"entrada": "08:00", "saida": "17:00"}])
    )
    lote_id = enviar_lote(cliente).json()["lote_id"]

    corpo = _laudo(cliente, lote_id).json()
    assert corpo["planilhas"] == []
    assert corpo["consolidado"] is None
    assert corpo["total_registros"] == 0
    assert cliente.get(f"/api/lotes/{lote_id}/downloads").json() == []


def test_reexportacao_substitui_arquivos(tmp_path) -> None:
    cliente, _ = montar_cliente(tmp_path, folha_com([dict(REGISTRO_COMPLETO)]))
    lote_id = enviar_lote(cliente).json()["lote_id"]

    primeiro = _laudo(cliente, lote_id).json()
    segundo = _laudo(cliente, lote_id).json()

    assert len(segundo["planilhas"]) == 1
    assert segundo["consolidado"] is not None
    # não acumula arquivos no lote
    assert len(cliente.get(f"/api/lotes/{lote_id}/downloads").json()) == 2
    # o primeiro arquivo foi substituído
    assert cliente.get(primeiro["planilhas"][0]["url"]).status_code == 404


def test_exportar_lote_inexistente(tmp_path) -> None:
    cliente, _ = montar_cliente(tmp_path, folha_com([dict(REGISTRO_COMPLETO)]))
    assert cliente.post(f"/api/lotes/{'0' * 32}/exportar").status_code == 404


def test_download_arquivo_inexistente(tmp_path) -> None:
    cliente, _ = montar_cliente(tmp_path, folha_com([dict(REGISTRO_COMPLETO)]))
    lote_id = enviar_lote(cliente).json()["lote_id"]
    _laudo(cliente, lote_id)
    assert cliente.get(f"/api/lotes/{lote_id}/downloads/naoexiste").status_code == 404


def test_download_bloqueia_arquivo_de_outro_lote(tmp_path) -> None:
    cliente, _ = montar_cliente(tmp_path, folha_com([dict(REGISTRO_COMPLETO)]))
    lote_a = enviar_lote(cliente).json()["lote_id"]
    lote_b = enviar_lote(cliente).json()["lote_id"]

    arquivo = _laudo(cliente, lote_a).json()["planilhas"][0]
    arquivo_id = arquivo["url"].rsplit("/", 1)[-1]

    # o arquivo pertence ao lote A; pedir pelo lote B deve falhar
    assert cliente.get(f"/api/lotes/{lote_b}/downloads/{arquivo_id}").status_code == 404
